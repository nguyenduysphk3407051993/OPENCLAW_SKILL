# -*- coding: utf-8 -*-
"""bang_ma_tran.json  ->  bang_ma_tran.docx

Giữ nguyên form mẫu: 4 dòng tiêu đề gộp ô (Mức độ đánh giá / TNKQ / Tự luận /
Tổng / Tỷ lệ % điểm), thân bảng gộp dọc theo chủ đề, 3 dòng chân bảng
(Tổng số câu/lệnh hỏi - Tổng số điểm - Tỷ lệ %).

Dùng:  python ma_tran_to_docx.py <bang_ma_tran.json> [output.docx]
"""

from __future__ import annotations

import sys

from docx.enum.table import WD_TABLE_ALIGNMENT

from docx_common import (TABLE_W_CM, apply_widths, fmt_diem, fmt_int, fmt_pct,
                         get_muc_do, get_nhom, load_json, merge, new_doc, para,
                         repeat_header, save_doc, set_borders, set_cell_margins,
                         set_fixed_layout, set_row_height, shade, title, write)

# Console Windows mặc định là cp1252/cp437 -> in tiếng Việt sẽ ném
# UnicodeEncodeError. Ép UTF-8 để chạy giống nhau trên Windows lẫn Linux.
for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass

W_TT, W_CHUDE, W_NOIDUNG = 0.9, 2.4, 3.1


# --------------------------------------------------------------- tính toán
def diem_o(nhom, nd, g_ma, m_ma):
    """Điểm của 1 ô (đơn vị kiến thức x nhóm x mức độ)."""
    ex = (nd.get("diem") or {}).get(g_ma, {}).get(m_ma)
    if ex is not None:
        return float(ex)
    n = (nd.get("so_lenh_hoi") or {}).get(g_ma, {}).get(m_ma, 0) or 0
    dpl = nhom.get("diem_moi_lenh")
    return float(n) * float(dpl) if dpl else 0.0


def so_o(nd, g_ma, m_ma):
    return int((nd.get("so_lenh_hoi") or {}).get(g_ma, {}).get(m_ma, 0) or 0)


def iter_noi_dung(cfg):
    for cd in cfg["chu_de"]:
        for nd in cd["noi_dung"]:
            yield cd, nd


# ------------------------------------------------------------------ dựng
def build(cfg, out_path):
    muc_do = get_muc_do(cfg)
    nhom, parents = get_nhom(cfg)
    L, G = len(muc_do), len(nhom)
    tong_diem = float(cfg["cau_hinh"].get("tong_diem", 10))

    c0 = 3                       # cột số liệu đầu tiên
    cT = c0 + G * L              # cột "Tổng" đầu tiên
    cP = cT + L                  # cột "Tỷ lệ % điểm"
    NCOL = cP + 1
    n_body = sum(len(cd["noi_dung"]) for cd in cfg["chu_de"])
    NROW = 4 + n_body + 3
    f0 = 4 + n_body

    doc = new_doc(landscape=True)
    title(doc, cfg.get("tieu_de", "MA TRẬN ĐỀ KIỂM TRA ĐỊNH KÌ"))
    ti = cfg.get("thong_tin") or {}
    if ti:
        para(doc, " – ".join(x for x in [
            f"Môn: {ti['mon']}" if ti.get("mon") else None,
            f"Lớp: {ti['lop']}" if ti.get("lop") else None,
            f"Thời gian: {ti['thoi_gian']}" if ti.get("thoi_gian") else None,
            ti.get("hinh_thuc"),
        ] if x), italic=True, align="center", size=10, space_after=4)

    t = doc.add_table(rows=NROW, cols=NCOL)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_borders(t)
    set_fixed_layout(t)
    set_cell_margins(t, left=25, right=25)
    rest = (TABLE_W_CM - W_TT - W_CHUDE - W_NOIDUNG) / (G * L + L + 1)
    apply_widths(t, [W_TT, W_CHUDE, W_NOIDUNG] + [rest] * (G * L + L + 1))

    # ---------------- 4 dòng tiêu đề
    write(merge(t, 0, 0, 3, 0), "TT", bold=True, size=9)
    write(merge(t, 0, 1, 3, 1), "Chủ đề/\nChương", bold=True, size=9)
    write(merge(t, 0, 2, 3, 2), "Nội dung/đơn vị\nkiến thức", bold=True, size=9)
    write(merge(t, 0, c0, 0, cT - 1), "Mức độ đánh giá", bold=True, size=9)
    write(merge(t, 0, cT, 2, cP - 1), "Tổng", bold=True, size=9)
    write(merge(t, 0, cP, 3, cP), "Tỷ lệ\n% điểm", bold=True, size=9)

    ci = c0
    for p in parents:
        span = len(p["con"]) * L
        one = len(p["con"]) == 1 and p["con"][0]["ten"] == p["ten"]
        if one:                                   # ví dụ: "Tự luận"
            write(merge(t, 1, ci, 2, ci + span - 1), p["ten"], bold=True, size=9)
        else:
            write(merge(t, 1, ci, 1, ci + span - 1), p["ten"], bold=True, size=9)
            cj = ci
            for g in p["con"]:
                write(merge(t, 2, cj, 2, cj + L - 1), g["ten"],
                      italic=True, size=8.5)
                cj += L
        ci += span

    for c in range(c0, cP):
        write(t.cell(3, c), muc_do[(c - c0) % L]["ten_ngan"], bold=True, size=8.5)
    for r in range(4):
        repeat_header(t.rows[r])
    set_row_height(t.rows[3], 0.85)

    # ---------------- thân bảng
    r = 4
    for cd in cfg["chu_de"]:
        top, bot = r, r + len(cd["noi_dung"]) - 1
        write(merge(t, top, 0, bot, 0), cd.get("tt", ""), size=9)
        write(merge(t, top, 1, bot, 1), cd["ten"], size=9, align="left")
        for nd in cd["noi_dung"]:
            write(t.cell(r, 2), nd["ten"], size=8.5, align="left")
            diem_dong = 0.0
            for mi, m in enumerate(muc_do):
                tong_m = 0
                for gi, g in enumerate(nhom):
                    n = so_o(nd, g["ma"], m["ma"])
                    write(t.cell(r, c0 + gi * L + mi), fmt_int(n), size=9)
                    tong_m += n
                    diem_dong += diem_o(g, nd, g["ma"], m["ma"])
                write(t.cell(r, cT + mi), fmt_int(tong_m), size=9, bold=True)
            write(t.cell(r, cP), fmt_pct(diem_dong / tong_diem * 100), size=9)
            set_row_height(t.rows[r], 0.55)
            r += 1

    # ---------------- 3 dòng chân bảng
    for k, lab in enumerate(("Tổng số câu/lệnh hỏi", "Tổng số điểm", "Tỷ lệ %")):
        c = merge(t, f0 + k, 0, f0 + k, 2)
        write(c, lab, bold=True, align="right", size=9)
        shade(c)
        set_row_height(t.rows[f0 + k], 0.5)

    # dòng 1: số câu/lệnh hỏi từng ô
    tong_g_diem = [0.0] * G
    tong_m_diem = [0.0] * L
    for gi, g in enumerate(nhom):
        for mi, m in enumerate(muc_do):
            n = sum(so_o(nd, g["ma"], m["ma"]) for _, nd in iter_noi_dung(cfg))
            d = sum(diem_o(g, nd, g["ma"], m["ma"]) for _, nd in iter_noi_dung(cfg))
            write(t.cell(f0, c0 + gi * L + mi), fmt_int(n), bold=True, size=9)
            tong_g_diem[gi] += d
            tong_m_diem[mi] += d
    for mi, m in enumerate(muc_do):
        n = sum(so_o(nd, g["ma"], m["ma"])
                for g in nhom for _, nd in iter_noi_dung(cfg))
        write(t.cell(f0, cT + mi), fmt_int(n), bold=True, size=9)

    # dòng 2 + 3: điểm và tỷ lệ theo nhóm (gộp ngang), theo mức độ (từng cột)
    for gi in range(G):
        a, b = c0 + gi * L, c0 + gi * L + L - 1
        write(merge(t, f0 + 1, a, f0 + 1, b), fmt_diem(tong_g_diem[gi]),
              bold=True, size=9)
        write(merge(t, f0 + 2, a, f0 + 2, b),
              fmt_pct(tong_g_diem[gi] / tong_diem * 100), bold=True, size=9)
    for mi in range(L):
        write(t.cell(f0 + 1, cT + mi), fmt_diem(tong_m_diem[mi]), bold=True, size=9)
        write(t.cell(f0 + 2, cT + mi),
              fmt_pct(tong_m_diem[mi] / tong_diem * 100), bold=True, size=9)
    write(t.cell(f0 + 1, cP), fmt_diem(sum(tong_m_diem)), bold=True, size=9)
    write(t.cell(f0 + 2, cP), fmt_pct(100), bold=True, size=9)

    for gc in cfg.get("ghi_chu", []):
        para(doc, gc, italic=True, size=9, space_before=2)

    save_doc(doc, out_path)
    return {"rows": NROW, "cols": NCOL,
            "tong_diem": sum(tong_m_diem),
            "diem_theo_muc_do": {muc_do[i]["ten"]: tong_m_diem[i] for i in range(L)},
            "diem_theo_nhom": {nhom[i]["ten"]: tong_g_diem[i] for i in range(G)}}


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else "bang_ma_tran.json"
    out = sys.argv[2] if len(sys.argv) > 2 else src.replace(".json", ".docx")
    info = build(load_json(src), out)
    print(f"[OK] {out}  ({info['rows']} dong x {info['cols']} cot)")
    print("     tong diem:", info["tong_diem"])


if __name__ == "__main__":
    main()
