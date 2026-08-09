# -*- coding: utf-8 -*-
"""bang_dac_ta_ma_tran.json  ->  bang_dac_ta_ma_tran.docx

Giữ nguyên form mẫu: 4 dòng tiêu đề (Số câu hỏi ở các mức độ đánh giá /
TNKQ / Tự luận / Biết-Hiểu-Vận dụng), cột "Yêu cầu cần đạt" liệt kê theo
từng mức độ, 3 dòng chân bảng.

Dùng:  python dac_ta_to_docx.py <bang_dac_ta_ma_tran.json> [output.docx]
"""

from __future__ import annotations

import sys

from docx.enum.table import WD_TABLE_ALIGNMENT

from docx_common import (TABLE_W_CM, apply_widths, fmt_diem, fmt_int, fmt_pct,
                         get_muc_do, get_nhom, load_json, merge, new_doc, para,
                         repeat_header, save_doc, set_borders, set_cell_margins,
                         set_fixed_layout, set_row_height, shade, title, write)

W_TT, W_CHUDE, W_NOIDUNG, W_YCCD = 0.9, 2.3, 2.8, 6.4


def _rows_of_nd(nd, muc_do):
    """Sinh danh sách dòng của 1 đơn vị kiến thức.

    Mỗi phần tử: (kieu, text, yccd) với kieu in {"label", "yccd"}.
    """
    rows = []
    for m in muc_do:
        items = [y for y in nd.get("yeu_cau_can_dat", [])
                 if y.get("muc_do") == m["ma"]]
        if not items:
            continue
        rows.append(("label", m["ten"] + ":", None))
        for y in items:
            rows.append(("yccd", y["noi_dung"], y))
    if not rows:
        rows.append(("label", "", None))
    return rows


def build(cfg, out_path):
    muc_do = get_muc_do(cfg)
    nhom, parents = get_nhom(cfg)
    L, G = len(muc_do), len(nhom)
    tong_diem = float(cfg["cau_hinh"].get("tong_diem", 10))

    c0 = 4
    NCOL = c0 + G * L

    plan = []           # (chu_de, [(noi_dung, rows)])
    n_body = 0
    for cd in cfg["chu_de"]:
        nds = []
        for nd in cd["noi_dung"]:
            rr = _rows_of_nd(nd, muc_do)
            nds.append((nd, rr))
            n_body += len(rr)
        plan.append((cd, nds))
    NROW = 4 + n_body + 3
    f0 = 4 + n_body

    doc = new_doc(landscape=True)
    title(doc, cfg.get("tieu_de", "BẢNG ĐẶC TẢ ĐỀ KIỂM TRA ĐỊNH KÌ"))
    ti = cfg.get("thong_tin") or {}
    if ti:
        para(doc, " – ".join(x for x in [
            f"Môn: {ti['mon']}" if ti.get("mon") else None,
            f"Lớp: {ti['lop']}" if ti.get("lop") else None,
            f"Thời gian: {ti['thoi_gian']}" if ti.get("thoi_gian") else None,
            ti.get("hinh_thuc"),
        ] if x), italic=True, align="center", size=10, space_after=4)

    t = doc.add_table(rows=NROW, cols=NCOL)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_borders(t)
    set_fixed_layout(t)
    set_cell_margins(t, left=25, right=25)
    rest = (TABLE_W_CM - W_TT - W_CHUDE - W_NOIDUNG - W_YCCD) / (G * L)
    apply_widths(t, [W_TT, W_CHUDE, W_NOIDUNG, W_YCCD] + [rest] * (G * L))

    # ---------------- 4 dòng tiêu đề
    write(merge(t, 0, 0, 3, 0), "TT", bold=True, size=9)
    write(merge(t, 0, 1, 3, 1), "Chủ đề/\nChương", bold=True, size=9)
    write(merge(t, 0, 2, 3, 2), "Nội dung/đơn vị\nkiến thức", bold=True, size=9)
    write(merge(t, 0, 3, 3, 3), "Yêu cầu cần đạt", bold=True, size=9)
    write(merge(t, 0, c0, 0, NCOL - 1),
          "Số câu hỏi ở các mức độ đánh giá", bold=True, size=9)

    ci = c0
    for p in parents:
        span = len(p["con"]) * L
        one = len(p["con"]) == 1 and p["con"][0]["ten"] == p["ten"]
        if one:
            write(merge(t, 1, ci, 2, ci + span - 1), p["ten"], bold=True, size=9)
        else:
            write(merge(t, 1, ci, 1, ci + span - 1), p["ten"], bold=True, size=9)
            cj = ci
            for g in p["con"]:
                write(merge(t, 2, cj, 2, cj + L - 1), g["ten"],
                      italic=True, size=8.5)
                cj += L
        ci += span

    for c in range(c0, NCOL):
        write(t.cell(3, c), muc_do[(c - c0) % L]["ten_ngan"], bold=True, size=8.5)
    for r in range(4):
        repeat_header(t.rows[r])
    set_row_height(t.rows[3], 0.85)

    # ---------------- thân bảng
    r = 4
    for cd, nds in plan:
        n_cd = sum(len(rr) for _, rr in nds)
        write(merge(t, r, 0, r + n_cd - 1, 0), cd.get("tt", ""), size=9)
        write(merge(t, r, 1, r + n_cd - 1, 1), cd["ten"], size=9, align="left")
        for nd, rr in nds:
            write(merge(t, r, 2, r + len(rr) - 1, 2), nd["ten"],
                  size=8.5, align="left")
            for kieu, text, y in rr:
                write(t.cell(r, 3), text, bold=(kieu == "label"),
                      align="left", size=8.5, valign="center")
                if y:
                    for g_ma, n in (y.get("so_lenh_hoi") or {}).items():
                        gi = next((i for i, g in enumerate(nhom)
                                   if g["ma"] == g_ma), None)
                        mi = next((i for i, m in enumerate(muc_do)
                                   if m["ma"] == y["muc_do"]), None)
                        if gi is None or mi is None or not n:
                            continue
                        write(t.cell(r, c0 + gi * L + mi), fmt_int(n),
                              bold=True, size=9)
                set_row_height(t.rows[r], 0.48)
                r += 1

    # ---------------- 3 dòng chân bảng
    for k, lab in enumerate(("Tổng số câu/lệnh hỏi", "Tổng số điểm", "Tỷ lệ %")):
        c = merge(t, f0 + k, 0, f0 + k, 2)
        write(c, lab, bold=True, align="right", size=9)
        shade(c)
        shade(t.cell(f0 + k, 3))
        set_row_height(t.rows[f0 + k], 0.5)

    tong_g_diem = [0.0] * G
    for gi, g in enumerate(nhom):
        dpl = g.get("diem_moi_lenh")
        for mi, m in enumerate(muc_do):
            n = 0
            d = 0.0
            for cd, nds in plan:
                for nd, _ in nds:
                    for y in nd.get("yeu_cau_can_dat", []):
                        if y.get("muc_do") != m["ma"]:
                            continue
                        k = (y.get("so_lenh_hoi") or {}).get(g["ma"], 0) or 0
                        n += k
                        ex = (y.get("diem") or {}).get(g["ma"])
                        d += float(ex) if ex is not None else k * float(dpl or 0)
            write(t.cell(f0, c0 + gi * L + mi), fmt_int(n), bold=True, size=9)
            tong_g_diem[gi] += d
        a, b = c0 + gi * L, c0 + gi * L + L - 1
        write(merge(t, f0 + 1, a, f0 + 1, b), fmt_diem(tong_g_diem[gi]),
              bold=True, size=9)
        write(merge(t, f0 + 2, a, f0 + 2, b),
              fmt_pct(tong_g_diem[gi] / tong_diem * 100), bold=True, size=9)

    for gc in cfg.get("ghi_chu", []):
        para(doc, gc, italic=True, size=9, space_before=2)

    save_doc(doc, out_path)
    return {"rows": NROW, "cols": NCOL, "tong_diem": sum(tong_g_diem)}


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else "bang_dac_ta_ma_tran.json"
    out = sys.argv[2] if len(sys.argv) > 2 else src.replace(".json", ".docx")
    info = build(load_json(src), out)
    print(f"[OK] {out}  ({info['rows']} dong x {info['cols']} cot)")
    print("     tong diem:", info["tong_diem"])


if __name__ == "__main__":
    main()
